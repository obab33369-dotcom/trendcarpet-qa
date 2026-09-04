import http.server
import socketserver
import os
import json
import urllib.parse
import sys
import subprocess
import mimetypes
import re
import datetime
import openpyxl
import hashlib
from PIL import Image

PORT = 8092
WORKSPACE_DIR = os.path.dirname(os.path.abspath(__file__))
BATCHES_FILE = os.path.join(WORKSPACE_DIR, "batches_config.json")
CALENDAR_CACHE_FILE = os.path.join(WORKSPACE_DIR, "calendar_projects.json")
THUMB_CACHE_DIR = os.path.join(WORKSPACE_DIR, ".cache", "thumbnails")
os.makedirs(THUMB_CACHE_DIR, exist_ok=True)
EXCEL_MASTER_FILE = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Fotostudion 230530-ongoing.xlsx"
EXCEL_2026_FILE = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\Projektlista_2026.xlsx"
BASE_SP_URL = "https://trendcarpet-my.sharepoint.com"

def clean_sp_url(target):
    if not target:
        return ""
    if target.startswith("http"):
        return target
    target = target.replace("../../../", "/")
    if not target.startswith("/"):
        target = "/" + target
    return BASE_SP_URL + target

def parse_project_date(d_str):
    if not d_str:
        return ("", "", "")
    d_str = str(d_str).split(" ")[0].strip()
    m = re.search(r'(\d{4})[-_ ](\d{1,2})[-_ ](\d{1,2})', d_str)
    if m:
        y, mo, d = m.group(1), m.group(2).zfill(2), m.group(3).zfill(2)
        return (f"{y}-{mo}-{d}", f"{y}-{mo}", y)
    m2 = re.search(r'(\d{2})[-_ ](\d{1,2})[-_ ](\d{1,2})', d_str)
    if m2:
        y, mo, d = "20" + m2.group(1), m2.group(2).zfill(2), m2.group(3).zfill(2)
        return (f"{y}-{mo}-{d}", f"{y}-{mo}", y)
    return (d_str, "", "")

def load_calendar_projects(force_refresh=False):
    if not force_refresh and os.path.exists(CALENDAR_CACHE_FILE):
        try:
            with open(CALENDAR_CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass

    all_projects = []

    # 1. Parse Fotostudion 230530-ongoing.xlsx
    if os.path.exists(EXCEL_MASTER_FILE):
        try:
            wb = openpyxl.load_workbook(EXCEL_MASTER_FILE, data_only=False)
            for sname in wb.sheetnames:
                ws = wb[sname]
                for r in range(2, ws.max_row + 1):
                    c_date = ws.cell(r, 1)
                    c_proj = ws.cell(r, 2)
                    c_photo = ws.cell(r, 3)
                    c_interiors = ws.cell(r, 4)
                    c_uploaded = ws.cell(r, 5)
                    c_png = ws.cell(r, 6)
                    c_orig = ws.cell(r, 7) if ws.max_column >= 7 else None

                    date_raw = str(c_date.value or "").strip()
                    proj_title = str(c_proj.value or "").strip()
                    if not proj_title and not date_raw:
                        continue

                    iso_date, year_month, year = parse_project_date(date_raw)
                    proj_link = clean_sp_url(c_proj.hyperlink.target if c_proj.hyperlink else "")
                    png_link = clean_sp_url(c_png.hyperlink.target if c_png and c_png.hyperlink else "")
                    orig_link = clean_sp_url(c_orig.hyperlink.target if c_orig and c_orig.hyperlink else "")

                    all_projects.append({
                        "id": f"fs_{sname}_{r}",
                        "sheet": sname,
                        "raw_date": date_raw,
                        "date": iso_date,
                        "year_month": year_month,
                        "year": year,
                        "title": proj_title,
                        "link": proj_link,
                        "photo_status": str(c_photo.value or "").strip(),
                        "interiors": str(c_interiors.value or "").strip(),
                        "uploaded": str(c_uploaded.value or "").strip(),
                        "png_link": png_link,
                        "orig_link": orig_link
                    })
        except Exception as e:
            print(f"Error parsing master Excel: {e}")

    # 2. Parse Projektlista_2026.xlsx
    if os.path.exists(EXCEL_2026_FILE):
        try:
            wb2 = openpyxl.load_workbook(EXCEL_2026_FILE, data_only=False)
            ws2 = wb2.active
            for r in range(2, ws2.max_row + 1):
                c_date = ws2.cell(r, 1)
                c_proj = ws2.cell(r, 2)
                c_photo = ws2.cell(r, 3)
                c_interiors = ws2.cell(r, 4)
                c_uploaded = ws2.cell(r, 5)

                date_raw = str(c_date.value or "").strip()
                proj_title = str(c_proj.value or "").strip()
                if not proj_title or date_raw.startswith("1905"):
                    continue

                iso_date, year_month, year = parse_project_date(date_raw)
                proj_link = clean_sp_url(c_proj.hyperlink.target if c_proj.hyperlink else "")

                if not any(p["title"].lower() == proj_title.lower() for p in all_projects):
                    all_projects.append({
                        "id": f"p2026_{r}",
                        "sheet": "Projekt 2026",
                        "raw_date": date_raw,
                        "date": iso_date,
                        "year_month": year_month,
                        "year": year,
                        "title": proj_title,
                        "link": proj_link,
                        "photo_status": str(c_photo.value or "").strip(),
                        "interiors": str(c_interiors.value or "").strip(),
                        "uploaded": str(c_uploaded.value or "").strip(),
                        "png_link": "",
                        "orig_link": ""
                    })
        except Exception as e:
            print(f"Error parsing 2026 Excel: {e}")

    try:
        with open(CALENDAR_CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(all_projects, f, indent=2, ensure_ascii=False)
    except Exception as e:
        print(f"Error caching calendar projects: {e}")

    return all_projects

def get_onedrive_web_url_for_path(local_path):
    if not local_path:
        return ""
    base_od = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB"
    norm = os.path.normpath(local_path)
    if norm.lower().startswith(base_od.lower()):
        rel = os.path.relpath(norm, base_od).replace('\\', '/')
        encoded_rel = urllib.parse.quote(f"/personal/andronik_lindgren_trendcarpet_se/Documents/{rel}", safe='')
        return f"https://trendcarpet-my.sharepoint.com/my?id={encoded_rel}"
    return ""

def load_batches():
    batches = []
    if os.path.exists(BATCHES_FILE):
        try:
            with open(BATCHES_FILE, "r", encoding="utf-8") as f:
                batches = json.load(f)
        except Exception:
            pass
    if not batches:
        batches = [
            {
                "id": "26-08-28-hatshop-black-river",
                "name": "2026-08-28 Hatshop Black River (14 modeller)",
                "path": r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\26-08-28-Hatshop-Black-River pt1-x",
                "category": "kepsar",
                "onedrive_url": "",
                "createdAt": "2026-09-02T11:00:00Z"
            }
        ]
        save_batches(batches)

    # Auto-match with calendar_projects if onedrive_url is missing
    changed = False
    try:
        projects = load_calendar_projects()
        for b in batches:
            if not b.get("onedrive_url"):
                folder = os.path.basename(os.path.normpath(b.get("path", ""))).strip().lower()
                bname = b.get("name", "").lower()
                matched_link = ""
                for p in projects:
                    ptitle = p.get("title", "").strip().lower()
                    if ptitle and p.get("link"):
                        if folder == ptitle or folder in ptitle or ptitle in folder:
                            matched_link = p["link"]
                            break
                if matched_link:
                    b["onedrive_url"] = matched_link
                    changed = True
                else:
                    auto_url = get_onedrive_web_url_for_path(b.get("path", ""))
                    if auto_url:
                        b["onedrive_url"] = auto_url
                        changed = True
        if changed:
            save_batches(batches)
    except Exception:
        pass

    return batches

def save_batches(batches):
    try:
        with open(BATCHES_FILE, "w", encoding="utf-8") as f:
            json.dump(batches, f, indent=2, ensure_ascii=False)
        return True
    except Exception as e:
        print(f"Error saving batches: {e}")
        return False

def get_batch(batch_id):
    batches = load_batches()
    if not batch_id:
        return batches[0] if batches else None
    for b in batches:
        if b["id"].lower() == batch_id.lower():
            return b
    return batches[0] if batches else None

def get_review_file(batch_id):
    safe_id = re.sub(r'[^a-zA-Z0-9_\-]', '_', batch_id)
    return os.path.join(WORKSPACE_DIR, f"reviews_{safe_id}.json")

def load_reviews(batch_id):
    rf = get_review_file(batch_id)
    if os.path.exists(rf):
        try:
            with open(rf, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    legacy = os.path.join(WORKSPACE_DIR, "hatshop_review_status.json")
    if batch_id == "26-08-28-hatshop-black-river" and os.path.exists(legacy):
        try:
            with open(legacy, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_reviews(batch_id, data):
    rf = get_review_file(batch_id)
    try:
        with open(rf, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        return True
    except Exception as e:
        print(f"Error saving reviews: {e}")
        return False

def parse_file_details(filename, folder_path):
    """
    Parses a filename to extract base shot number, version suffix, modification time, etc.
    Examples:
      keps-model-color-01.jpg -> base_shot='01', suffix='', is_version_suffix=False
      keps-model-color-01-n.jpg -> base_shot='01', suffix='n', is_version_suffix=True
      matta-model-01-Interior-0002.jpg -> base_shot='01', suffix='Interior-0002', is_version_suffix=False
    """
    name_no_ext, ext = os.path.splitext(filename)
    full_path = os.path.join(folder_path, filename)
    mtime = os.path.getmtime(full_path) if os.path.exists(full_path) else 0
    file_size = os.path.getsize(full_path) if os.path.exists(full_path) else 0

    m = re.search(r'[-_](0[1-9]|[1-9][0-9]?)(.*)$', name_no_ext, re.IGNORECASE)
    if m:
        base_shot_num = m.group(1).zfill(2)
        raw_suffix = m.group(2).strip('-_ ')
        clean_suffix = raw_suffix.replace('-', ' ').replace('_', ' ').strip()
        
        # Check if true replacement version (e.g. -n, -ny, -v2, -fix, -2)
        is_interior = bool(re.search(r'interioi?r|miljo|miljö|detail', raw_suffix, re.IGNORECASE))
        is_version_suffix = bool(re.match(r'^(n|ny|new|v[0-9]+|[0-9]+|fix|crop|justerad|redigerad)$', raw_suffix, re.IGNORECASE)) and not is_interior
        
        return {
            "filename": filename,
            "base_shot": base_shot_num,
            "raw_suffix": raw_suffix,
            "clean_suffix": clean_suffix,
            "is_version_suffix": is_version_suffix,
            "is_interior": is_interior,
            "mtime": mtime,
            "size_bytes": file_size
        }
    else:
        return {
            "filename": filename,
            "base_shot": "99",
            "raw_suffix": "",
            "clean_suffix": "",
            "is_version_suffix": False,
            "is_interior": False,
            "mtime": mtime,
            "size_bytes": file_size
        }

def parse_upload_order_file(txt_path):
    order = []
    if not os.path.exists(txt_path):
        return order
    try:
        with open(txt_path, 'r', encoding='utf-8', errors='ignore') as f:
            for line in f:
                line = line.strip()
                if '|' in line and not line.startswith('ORDNING') and not line.startswith('---'):
                    parts = [p.strip() for p in line.split('|')]
                    if len(parts) >= 2:
                        seq = parts[0].zfill(2)
                        new_fn = parts[1]
                        orig_fn = parts[2] if len(parts) > 2 else ''
                        typ = parts[3] if len(parts) > 3 else ''
                        order.append({
                            'seq': seq,
                            'filename': new_fn,
                            'orig_filename': orig_fn,
                            'type': typ
                        })
    except Exception as e:
        print(f"Error reading upload_order.txt {txt_path}: {e}")
    return order

def process_upload_order_shots(upload_order, raw_files, batch_id, subdir, folder_name, folder_path):
    final_shots = []
    seen_files = set()

    for item in upload_order:
        fn = item["filename"]
        full_fn = os.path.join(folder_path, fn)
        if not os.path.exists(full_fn):
            match = next((f for f in raw_files if f.lower() == fn.lower()), None)
            if match:
                fn = match
                full_fn = os.path.join(folder_path, fn)
            else:
                continue

        seen_files.add(fn.lower())
        mtime = os.path.getmtime(full_fn) if os.path.exists(full_fn) else 0
        file_size = os.path.getsize(full_fn) if os.path.exists(full_fn) else 0
        rel_path = f"{subdir}/{folder_name}/{fn}" if subdir else f"{folder_name}/{fn}"
        raw_url = f"/image/{urllib.parse.quote(batch_id)}/{urllib.parse.quote(rel_path)}?v={int(mtime)}"

        typ_upper = item["type"].upper()
        is_bubble = "BUBBLA" in typ_upper
        is_gallery = "GALLERI" in typ_upper

        if is_bubble:
            label = f"{item['seq']} (Bubbla)"
            v_tag = "Bubbla"
            is_new = True
        elif is_gallery:
            if "interior" in fn.lower() or "miljo" in fn.lower():
                label = f"{item['seq']} (Miljö)"
            elif "-01" in fn.lower():
                label = f"{item['seq']} (Plan)"
            else:
                label = f"{item['seq']}"
            v_tag = ""
            is_new = False
        else:
            label = f"{item['seq']}"
            v_tag = ""
            is_new = False

        final_shots.append({
            "filename": fn,
            "shot": item["seq"],
            "base_shot": item["seq"],
            "label": label,
            "is_primary": (item["seq"] == "01"),
            "is_new_version": is_new,
            "version_tag": v_tag,
            "mtime": mtime,
            "url": raw_url,
            "thumb_url": f"{raw_url}&thumb=160",
            "rel_path": rel_path,
            "size_bytes": file_size
        })

    # Remaining files if any
    for f in raw_files:
        if f.lower() not in seen_files:
            full_fn = os.path.join(folder_path, f)
            mtime = os.path.getmtime(full_fn) if os.path.exists(full_fn) else 0
            file_size = os.path.getsize(full_fn) if os.path.exists(full_fn) else 0
            rel_path = f"{subdir}/{folder_name}/{f}" if subdir else f"{folder_name}/{f}"
            raw_url = f"/image/{urllib.parse.quote(batch_id)}/{urllib.parse.quote(rel_path)}?v={int(mtime)}"
            final_shots.append({
                "filename": f,
                "shot": "99",
                "base_shot": "99",
                "label": f,
                "is_primary": False,
                "is_new_version": False,
                "version_tag": "",
                "mtime": mtime,
                "url": raw_url,
                "thumb_url": f"{raw_url}&thumb=160",
                "rel_path": rel_path,
                "size_bytes": file_size
            })

    return final_shots

def process_and_sort_variant_shots(raw_files, batch_id, subdir, folder_name, folder_path):
    """
    Processes all raw image files for a product variant.
    If upload_order.txt exists, respects the explicit studio sequence and types.
    Otherwise groups versioned replacement shots (e.g. 01.jpg and 01-n.jpg) by regex.
    """
    upload_order_path = os.path.join(folder_path, "upload_order.txt")
    if os.path.exists(upload_order_path):
        parsed_order = parse_upload_order_file(upload_order_path)
        if parsed_order:
            return process_upload_order_shots(parsed_order, raw_files, batch_id, subdir, folder_name, folder_path)

    parsed_files = [parse_file_details(f, folder_path) for f in raw_files]
    
    version_groups = {}
    independent_shots = []

    for p in parsed_files:
        if p["is_version_suffix"] or not p["raw_suffix"]:
            version_groups.setdefault(p["base_shot"], []).append(p)
        else:
            independent_shots.append(p)

    final_shots = []

    for base_shot in sorted(version_groups.keys()):
        group = version_groups[base_shot]
        if len(group) == 1:
            item = group[0]
            rel_path = f"{subdir}/{folder_name}/{item['filename']}" if subdir else f"{folder_name}/{item['filename']}"
            has_suffix = bool(item["raw_suffix"])
            label = f"{base_shot} ({item['clean_suffix']})" if has_suffix else f"{base_shot}"
            raw_url = f"/image/{urllib.parse.quote(batch_id)}/{urllib.parse.quote(rel_path)}?v={int(item['mtime'])}"
            
            final_shots.append({
                "filename": item["filename"],
                "shot": base_shot,
                "base_shot": base_shot,
                "label": label,
                "is_primary": True,
                "is_new_version": item["is_version_suffix"],
                "version_tag": f"-{item['raw_suffix']}" if item["raw_suffix"] else "",
                "mtime": item["mtime"],
                "url": raw_url,
                "thumb_url": f"{raw_url}&thumb=160",
                "rel_path": rel_path,
                "size_bytes": item["size_bytes"]
            })
        else:
            group.sort(key=lambda x: (1 if x["is_version_suffix"] else 0, x["mtime"]), reverse=True)
            primary = group[0]
            rel_path_prim = f"{subdir}/{folder_name}/{primary['filename']}" if subdir else f"{folder_name}/{primary['filename']}"
            version_str = f"-{primary['raw_suffix']}" if primary['raw_suffix'] else "ny"
            label_prim = f"{base_shot} (Ny {version_str})" if primary['raw_suffix'] else f"{base_shot} (Ny)"
            raw_url_prim = f"/image/{urllib.parse.quote(batch_id)}/{urllib.parse.quote(rel_path_prim)}?v={int(primary['mtime'])}"

            final_shots.append({
                "filename": primary["filename"],
                "shot": base_shot,
                "base_shot": base_shot,
                "label": label_prim,
                "is_primary": True,
                "is_new_version": True,
                "version_tag": version_str,
                "has_previous_version": True,
                "mtime": primary["mtime"],
                "url": raw_url_prim,
                "thumb_url": f"{raw_url_prim}&thumb=160",
                "rel_path": rel_path_prim,
                "size_bytes": primary["size_bytes"]
            })

            for idx, older in enumerate(group[1:], start=1):
                rel_path_old = f"{subdir}/{folder_name}/{older['filename']}" if subdir else f"{folder_name}/{older['filename']}"
                old_suffix = f"-{older['raw_suffix']}" if older['raw_suffix'] else "orig"
                raw_url_old = f"/image/{urllib.parse.quote(batch_id)}/{urllib.parse.quote(rel_path_old)}?v={int(older['mtime'])}"
                final_shots.append({
                    "filename": older["filename"],
                    "shot": f"{base_shot}_old{idx}",
                    "base_shot": base_shot,
                    "label": f"{base_shot} (Tidigare {old_suffix})",
                    "is_primary": False,
                    "is_new_version": False,
                    "version_tag": "tidigare",
                    "mtime": older["mtime"],
                    "url": raw_url_old,
                    "thumb_url": f"{raw_url_old}&thumb=160",
                    "rel_path": rel_path_old,
                    "size_bytes": older["size_bytes"]
                })

    for item in independent_shots:
        rel_path = f"{subdir}/{folder_name}/{item['filename']}" if subdir else f"{folder_name}/{item['filename']}"
        clean_name = item["clean_suffix"]
        clean_name = re.sub(r'interioi?r', 'Miljö', clean_name, flags=re.IGNORECASE)
        label = f"{item['base_shot']} ({clean_name})"
        raw_url = f"/image/{urllib.parse.quote(batch_id)}/{urllib.parse.quote(rel_path)}?v={int(item['mtime'])}"
        final_shots.append({
            "filename": item["filename"],
            "shot": f"{item['base_shot']}_{item['raw_suffix']}",
            "base_shot": item["base_shot"],
            "label": label,
            "is_primary": True,
            "is_new_version": False,
            "version_tag": "",
            "mtime": item["mtime"],
            "url": raw_url,
            "thumb_url": f"{raw_url}&thumb=160",
            "rel_path": rel_path,
            "size_bytes": item["size_bytes"]
        })

    # Sort: base shot asc, studio planvy before miljö/sub-shots, is_primary desc
    final_shots.sort(key=lambda s: (s["base_shot"], 0 if not s["label"].endswith(")") else 1, 0 if s["is_primary"] else 1, s["filename"]))
    return final_shots

def scan_batch_products(batch):
    base_dir = batch.get("path", "")
    batch_id = batch.get("id", "")
    if not os.path.exists(base_dir):
        return []

    root_entries = [d for d in os.listdir(base_dir) if os.path.isdir(os.path.join(base_dir, d))]
    has_variant_subdirs = any(v.lower() in ["1500px", "1500px iphone", "original", "web", "print"] for v in root_entries)

    products_map = {}

    if has_variant_subdirs:
        for subdir in root_entries:
            subdir_path = os.path.join(base_dir, subdir)
            if not os.path.isdir(subdir_path):
                continue
            
            for folder_name in sorted(os.listdir(subdir_path)):
                folder_path = os.path.join(subdir_path, folder_name)
                if not os.path.isdir(folder_path):
                    continue

                parts = folder_name.split("-")
                brand = "Brand"
                model_name = folder_name
                color_name = ""

                if len(parts) >= 4 and parts[0].strip().lower() in ["keps", "matta", "wiltonmatta", "inredning"]:
                    brand = parts[0].strip()
                    model_name = parts[1].strip()
                    color_name = "-".join(parts[2:]).strip()
                elif len(parts) >= 3:
                    model_name = parts[1].strip()
                    color_name = "-".join(parts[2:]).strip()
                elif len(parts) == 2:
                    model_name = parts[0].strip()
                    color_name = parts[1].strip()

                if folder_name not in products_map:
                    products_map[folder_name] = {
                        "id": folder_name,
                        "brand": brand,
                        "model": model_name,
                        "color": color_name,
                        "folder_name": folder_name,
                        "has_new_version": False,
                        "version_tag": "",
                        "latest_mtime": 0,
                        "variants": {}
                    }

                raw_files = [f for f in os.listdir(folder_path) if f.lower().endswith((".jpg", ".jpeg", ".png", ".webp"))]
                shot_list = process_and_sort_variant_shots(raw_files, batch_id, subdir, folder_name, folder_path)

                products_map[folder_name]["variants"][subdir] = shot_list
                if any(s.get("is_new_version") for s in shot_list):
                    products_map[folder_name]["has_new_version"] = True
                    tag = next((s["version_tag"] for s in shot_list if s.get("is_new_version")), "")
                    if tag:
                        products_map[folder_name]["version_tag"] = tag
                
                mtimes = [s["mtime"] for s in shot_list if s.get("mtime")]
                if mtimes:
                    products_map[folder_name]["latest_mtime"] = max(products_map[folder_name]["latest_mtime"], max(mtimes))
    else:
        for folder_name in sorted(root_entries):
            folder_path = os.path.join(base_dir, folder_name)
            parts = folder_name.split("-")
            brand = "Brand"
            model_name = folder_name
            color_name = ""

            if len(parts) >= 3:
                model_name = parts[1].strip()
                color_name = "-".join(parts[2:]).strip()
            elif len(parts) == 2:
                model_name = parts[0].strip()
                color_name = parts[1].strip()

            raw_files = [f for f in os.listdir(folder_path) if f.lower().endswith((".jpg", ".jpeg", ".png", ".webp"))]
            shot_list = process_and_sort_variant_shots(raw_files, batch_id, "", folder_name, folder_path)

            has_new = any(s.get("is_new_version") for s in shot_list)
            v_tag = next((s["version_tag"] for s in shot_list if s.get("is_new_version")), "")
            mtimes = [s["mtime"] for s in shot_list if s.get("mtime")]

            products_map[folder_name] = {
                "id": folder_name,
                "brand": brand,
                "model": model_name,
                "color": color_name,
                "folder_name": folder_name,
                "has_new_version": has_new,
                "version_tag": v_tag,
                "latest_mtime": max(mtimes) if mtimes else 0,
                "variants": { "Standard": shot_list }
            }

    products = list(products_map.values())
    products.sort(key=lambda x: x["model"])
    return products

class HubHTTPRequestHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=WORKSPACE_DIR, **kwargs)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        query = urllib.parse.parse_qs(parsed.query)

        if path in ["/", "/index.html", "/ui", "/hatshop_ui.html"]:
            user_agent = self.headers.get("User-Agent", "").lower()
            is_mobile = any(m in user_agent for m in ["iphone", "android", "mobile", "ipod", "blackberry", "opera mini"])
            force_desktop = "desktop" in query and query["desktop"][0] in ["1", "true"]

            if is_mobile and not force_desktop:
                redirect_url = "/mobile"
                if parsed.query:
                    redirect_url += f"?{parsed.query}"
                self.send_response(302)
                self.send_header("Location", redirect_url)
                self.send_header("Cache-Control", "no-cache")
                self.end_headers()
                return

            html_path = os.path.join(WORKSPACE_DIR, "hatshop_ui.html")
            if os.path.exists(html_path):
                with open(html_path, "rb") as f:
                    content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(content)))
                self.send_header("Cache-Control", "no-cache")
                self.end_headers()
                self.wfile.write(content)
                return
            else:
                self.send_error(404, "hatshop_ui.html not found")
                return

        elif path in ["/m", "/mobile", "/hatshop_mobile.html"]:
            html_path = os.path.join(WORKSPACE_DIR, "hatshop_mobile.html")
            if os.path.exists(html_path):
                with open(html_path, "rb") as f:
                    content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(content)))
                self.send_header("Cache-Control", "no-cache")
                self.end_headers()
                self.wfile.write(content)
                return
            else:
                self.send_error(404, "hatshop_mobile.html not found")
                return

        elif path == "/api/batches":
            batches = load_batches()
            content = json.dumps({"batches": batches}, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(content)
            return

        elif path == "/api/calendar":
            force = query.get("refresh", ["false"])[0].lower() == "true"
            projects = load_calendar_projects(force_refresh=force)
            content = json.dumps({"projects": projects, "count": len(projects)}, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(content)
            return

        elif path == "/api/products":
            batch_id = query.get("batch", [""])[0]
            batches = load_batches()
            batch = next((b for b in batches if b["id"] == batch_id), batches[0] if batches else None)

            if not batch:
                self.send_error(404, "Batch not found")
                return

            products = scan_batch_products(batch)
            reviews = load_reviews(batch["id"])
            payload = {
                "batch": batch,
                "count": len(products),
                "products": products,
                "reviews": reviews
            }
            content = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(content)
            return

        elif path == "/api/reviews":
            batch_id = query.get("batch", [""])[0]
            if not batch_id:
                batches = load_batches()
                batch_id = batches[0]["id"] if batches else "default"
            reviews = load_reviews(batch_id)
            resp = json.dumps({"status": "ok", "reviews": reviews}, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(resp)
            return

        elif path == "/api/open-folder":
            batch_id = query.get("batch", [""])[0]
            subpath = query.get("path", [""])[0]
            batch = get_batch(batch_id) if batch_id else load_batches()[0]
            base_dir = batch["path"] if batch else WORKSPACE_DIR

            if subpath:
                full_path = os.path.normpath(os.path.join(base_dir, subpath))
            else:
                full_path = base_dir

            if os.path.exists(full_path):
                if os.path.isfile(full_path):
                    subprocess.Popen(f'explorer /select,"{full_path}"')
                else:
                    subprocess.Popen(f'explorer "{full_path}"')
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(b'{"status":"ok"}')
                return
            else:
                self.send_error(404, "Directory or file does not exist")
                return

        elif path.startswith("/image/"):
            parts = path[len("/image/"):].split("/", 1)
            if len(parts) < 2:
                self.send_error(400, "Invalid image URL")
                return
            
            batch_id = urllib.parse.unquote(parts[0])
            rel_path = urllib.parse.unquote(parts[1])
            batch = get_batch(batch_id)
            if not batch:
                self.send_error(404, "Batch not found")
                return

            base_dir = batch["path"]
            file_path = os.path.normpath(os.path.join(base_dir, rel_path))

            if not os.path.abspath(file_path).startswith(os.path.abspath(base_dir)):
                self.send_error(403, "Access denied")
                return

            if os.path.exists(file_path) and os.path.isfile(file_path):
                file_to_serve = file_path
                is_thumb = "thumb" in query or "w" in query

                if is_thumb:
                    try:
                        raw_w = query.get("w", query.get("thumb", ["160"]))[0]
                        target_w = int(raw_w) if str(raw_w).isdigit() else 160
                    except Exception:
                        target_w = 160
                    if target_w <= 0 or target_w > 800:
                        target_w = 160

                    try:
                        cache_key = hashlib.md5(f"{file_path}_{os.path.getmtime(file_path)}_{target_w}".encode("utf-8")).hexdigest() + ".jpg"
                        cache_file = os.path.join(THUMB_CACHE_DIR, cache_key)
                        if not os.path.exists(cache_file):
                            with Image.open(file_path) as img:
                                img = img.convert("RGB")
                                img.thumbnail((target_w, target_w), Image.Resampling.LANCZOS)
                                img.save(cache_file, "JPEG", quality=82, optimize=True)
                        file_to_serve = cache_file
                    except Exception as e:
                        file_to_serve = file_path

                mime_type, _ = mimetypes.guess_type(file_to_serve)
                if not mime_type:
                    mime_type = "image/jpeg"
                
                try:
                    file_size = os.path.getsize(file_to_serve)
                    with open(file_to_serve, "rb") as f:
                        self.send_response(200)
                        self.send_header("Content-Type", mime_type)
                        self.send_header("Content-Length", str(file_size))
                        self.send_header("Cache-Control", "public, max-age=86400")
                        self.send_header("Access-Control-Allow-Origin", "*")
                        self.end_headers()
                        while True:
                            chunk = f.read(65536)
                            if not chunk:
                                break
                            self.wfile.write(chunk)
                    return
                except (BrokenPipeError, ConnectionResetError):
                    return
                except Exception as e:
                    try:
                        self.send_error(500, f"Error reading file: {e}")
                    except Exception:
                        pass
                    return
            else:
                self.send_error(404, f"Image not found: {rel_path}")
                return

        else:
            return super().do_GET()

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        query = urllib.parse.parse_qs(parsed.query)

        if path == "/api/batches/add":
            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                body = json.loads(post_data.decode("utf-8"))
                folder_path = body.get("path", "").strip().strip('"').strip("'")
                name = body.get("name", "").strip()
                category = body.get("category", "kepsar")

                if not os.path.exists(folder_path):
                    self.send_response(400)
                    self.send_header("Content-Type", "application/json; charset=utf-8")
                    self.send_header("Access-Control-Allow-Origin", "*")
                    self.end_headers()
                    self.wfile.write(json.dumps({"status": "error", "message": f"Sökvägen existerar inte på datorn: {folder_path}"}).encode("utf-8"))
                    return

                folder_base = os.path.basename(os.path.normpath(folder_path))
                if not name:
                    name = folder_base

                # Auto-count models in folder if not specified
                try:
                    temp_batch = {"id": "temp", "path": folder_path}
                    prods = scan_batch_products(temp_batch)
                    if prods and not re.search(r'\(\d+\s*modell', name, re.IGNORECASE):
                        name = f"{name} ({len(prods)} modeller)"
                except Exception:
                    pass

                clean_id_base = re.sub(r'\s*\(\d+\s*modeller?\)', '', name, flags=re.IGNORECASE)
                batch_id = re.sub(r'[^a-zA-Z0-9_-]', '-', clean_id_base.lower()).strip('-')
                batches = load_batches()
                existing = next((b for b in batches if b["id"] == batch_id), None)
                if existing:
                    batch_id = f"{batch_id}-{len(batches)+1}"

                new_batch = {
                    "id": batch_id,
                    "name": name,
                    "path": folder_path,
                    "category": category,
                    "onedrive_url": body.get("onedrive_url", "").strip(),
                    "createdAt": datetime.datetime.now().isoformat()
                }
                batches.append(new_batch)
                save_batches(batches)

                resp = json.dumps({"status": "ok", "batch": new_batch, "batches": batches}).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"Error adding batch: {e}")
                return

        elif path == "/api/batch/set_onedrive":
            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                body = json.loads(post_data.decode("utf-8")) if content_length > 0 else {}
                batch_id = body.get("batch_id") or query.get("batch", [""])[0]
                url = body.get("onedrive_url", "").strip()
                batches = load_batches()
                target_batch = next((b for b in batches if b["id"] == batch_id), None)
                if target_batch:
                    target_batch["onedrive_url"] = url
                    save_batches(batches)
                    resp = json.dumps({"status": "ok", "batch": target_batch, "batches": batches}).encode("utf-8")
                else:
                    resp = json.dumps({"status": "error", "message": "Batch not found"}).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"Error updating onedrive url: {e}")
                return

        elif path == "/api/comment/add":
            batch_id = query.get("batch", [""])[0]
            if not batch_id:
                batches = load_batches()
                batch_id = batches[0]["id"] if batches else "default"

            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                body = json.loads(post_data.decode("utf-8"))
                pid = body.get("product_id")
                author = body.get("author", "Anonym")
                text = body.get("text", "")
                shots = body.get("shots", [])

                if not pid or not text:
                    self.send_error(400, "Missing product_id or text")
                    return

                if not shots:
                    at_matches = re.findall(r'@(0[1-9]|[1-9][0-9]?)', text)
                    if at_matches:
                        shots = sorted(list(set(m.zfill(2) for m in at_matches)))

                current_reviews = load_reviews(batch_id)
                if pid not in current_reviews:
                    current_reviews[pid] = { "status": "pending", "flags": {}, "comments": [] }

                if "comments" not in current_reviews[pid]:
                    current_reviews[pid]["comments"] = []

                drawings = body.get("drawings", [])

                new_comment = {
                    "id": f"c_{int(datetime.datetime.now().timestamp()*1000)}",
                    "author": author,
                    "text": text,
                    "shots": shots,
                    "drawings": drawings,
                    "time": datetime.datetime.now().strftime("%H:%M")
                }
                current_reviews[pid]["comments"].append(new_comment)
                current_reviews[pid]["reviewer"] = author
                current_reviews[pid]["updatedAt"] = datetime.datetime.now().isoformat()

                save_reviews(batch_id, current_reviews)
                resp = json.dumps({"status": "ok", "comment": new_comment, "reviews": current_reviews}).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"Comment error: {e}")
                return

        elif path == "/api/comment/delete":
            batch_id = query.get("batch", [""])[0]
            if not batch_id:
                batches = load_batches()
                batch_id = batches[0]["id"] if batches else "default"

            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                body = json.loads(post_data.decode("utf-8"))
                pid = body.get("product_id")
                comment_id = body.get("comment_id")

                current_reviews = load_reviews(batch_id)
                if pid in current_reviews and "comments" in current_reviews[pid]:
                    current_reviews[pid]["comments"] = [
                        c for c in current_reviews[pid]["comments"] if c.get("id") != comment_id
                    ]
                    save_reviews(batch_id, current_reviews)

                resp = json.dumps({"status": "ok", "reviews": current_reviews}).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"Delete comment error: {e}")
                return

        elif path == "/api/comment/react":
            batch_id = query.get("batch", [""])[0]
            if not batch_id:
                batches = load_batches()
                batch_id = batches[0]["id"] if batches else "default"

            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                body = json.loads(post_data.decode("utf-8"))
                pid = body.get("product_id")
                comment_id = body.get("comment_id")
                emoji = body.get("emoji", "👍")
                author = body.get("author", "Anonym")

                current_reviews = load_reviews(batch_id)
                if pid in current_reviews and "comments" in current_reviews[pid]:
                    for c in current_reviews[pid]["comments"]:
                        if c.get("id") == comment_id:
                            reactions = c.setdefault("reactions", {})
                            user_list = reactions.setdefault(emoji, [])
                            if author in user_list:
                                user_list.remove(author)
                                if not user_list:
                                    del reactions[emoji]
                            else:
                                user_list.append(author)
                            break
                    save_reviews(batch_id, current_reviews)

                resp = json.dumps({"status": "ok", "reviews": current_reviews}).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"React error: {e}")
                return

        elif path == "/api/review":
            batch_id = query.get("batch", [""])[0]
            if not batch_id:
                batches = load_batches()
                batch_id = batches[0]["id"] if batches else "default"

            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                review_data = json.loads(post_data.decode("utf-8"))
                current_reviews = load_reviews(batch_id)
                
                if "product_id" in review_data:
                    pid = review_data["product_id"]
                    new_data = review_data.get("data", {})
                    if pid in current_reviews and "comments" in current_reviews[pid] and "comments" not in new_data:
                        new_data["comments"] = current_reviews[pid]["comments"]
                    current_reviews[pid] = new_data
                elif "reviews" in review_data:
                    current_reviews = review_data["reviews"]

                save_reviews(batch_id, current_reviews)
                resp = json.dumps({"status": "ok", "reviews": current_reviews}).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"Invalid JSON or save error: {e}")
                return

        elif path == "/api/status/set":
            batch_id = query.get("batch", [""])[0]
            if not batch_id:
                batches = load_batches()
                batch_id = batches[0]["id"] if batches else "default"

            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                body = json.loads(post_data.decode("utf-8"))
                pid = body.get("product_id")
                status = body.get("status", "pending")
                reviewer = body.get("reviewer", "Anonym")

                if not pid:
                    self.send_error(400, "Missing product_id")
                    return

                current_reviews = load_reviews(batch_id)
                if pid not in current_reviews:
                    current_reviews[pid] = { "status": status, "flags": {}, "comments": [] }
                else:
                    current_reviews[pid]["status"] = status

                current_reviews[pid]["reviewer"] = reviewer
                current_reviews[pid]["updatedAt"] = datetime.datetime.now().isoformat()

                save_reviews(batch_id, current_reviews)
                resp = json.dumps({"status": "ok", "reviews": current_reviews}, ensure_ascii=False).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"Status error: {e}")
                return

        elif path == "/api/checklist/toggle":
            batch_id = query.get("batch", [""])[0]
            if not batch_id:
                batches = load_batches()
                batch_id = batches[0]["id"] if batches else "default"

            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                body = json.loads(post_data.decode("utf-8"))
                pid = body.get("product_id")
                flag = body.get("flag")
                value = body.get("value")
                reviewer = body.get("reviewer", "Anonym")

                if not pid or not flag:
                    self.send_error(400, "Missing product_id or flag")
                    return

                current_reviews = load_reviews(batch_id)
                if pid not in current_reviews:
                    current_reviews[pid] = { "status": "pending", "flags": {}, "comments": [] }

                if "flags" not in current_reviews[pid]:
                    current_reviews[pid]["flags"] = {}

                current_reviews[pid]["flags"][flag] = bool(value)
                current_reviews[pid]["reviewer"] = reviewer
                current_reviews[pid]["updatedAt"] = datetime.datetime.now().isoformat()

                save_reviews(batch_id, current_reviews)
                resp = json.dumps({"status": "ok", "reviews": current_reviews}, ensure_ascii=False).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.send_header("Access-Control-Allow-Origin", "*")
                self.end_headers()
                self.wfile.write(resp)
                return
            except Exception as e:
                self.send_error(400, f"Checklist error: {e}")
                return

        else:
            self.send_error(404, "Endpoint not found")

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

def start_server(port=PORT):
    socketserver.ThreadingTCPServer.allow_reuse_address = True
    actual_port = port
    while actual_port < port + 10:
        try:
            server = socketserver.ThreadingTCPServer(("", actual_port), HubHTTPRequestHandler)
            break
        except OSError:
            actual_port += 1

    print("============================================================")
    print(" QA Photo Review Hub Server")
    print(f" Access URL: http://localhost:{actual_port}")
    print("============================================================", flush=True)
    server.serve_forever()

if __name__ == "__main__":
    p = PORT
    if len(sys.argv) > 1:
        try:
            p = int(sys.argv[1])
        except ValueError:
            pass
    start_server(p)
