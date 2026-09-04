import openpyxl
import pandas as pd
import os

path = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\26-06-05-in-out\Bade_IK100101936_Foto.xlsx"
print("File exists:", os.path.exists(path))

try:
    df = pd.read_excel(path)
    print("Pandas read successfully:")
    print(df.head(20))
    print(df.columns)
except Exception as e:
    print("Pandas error:", e)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        print("Openpyxl read successfully. Sheets:", wb.sheetnames)
        sheet = wb.active
        for r in range(1, min(sheet.max_row + 1, 30)):
            row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
            print(f"Row {r}: {row_vals}")
    except Exception as e2:
        print("Openpyxl error:", e2)
