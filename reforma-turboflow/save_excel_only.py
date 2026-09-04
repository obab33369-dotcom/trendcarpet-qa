import os
import glob
import openpyxl
from openpyxl.styles import Font

project_dir = r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\Projects\REFORMA\reforma-turboflow"
output_excel = os.path.join(project_dir, "turboflow_batches_index.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Batches"

headers = ["Batch Number", "Prompts TXT File (Click to Open)", "Folder Path (Copy this)", "Unique Images Count"]
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = Font(bold=True)

batch_files = glob.glob(os.path.join(project_dir, "turboflow_tracking_log_full_catalog_batch*.csv"))
num_batches = len(batch_files)

for i in range(1, num_batches + 1):
    txt_file = os.path.join(project_dir, f"turboflow_ready_full_catalog_batch{i}.txt")
    batch_img_dir = os.path.join(project_dir, f"batch{i}_images")
    
    # Just list files in the folder to get the image count
    if os.path.exists(batch_img_dir):
        image_count = len([f for f in os.listdir(batch_img_dir) if os.path.isfile(os.path.join(batch_img_dir, f))])
    else:
        image_count = 0
        
    ws.cell(row=i+1, column=1, value=f"Batch {i}")
    
    txt_cell = ws.cell(row=i+1, column=2)
    txt_cell.value = "Öppna TXT-fil"
    txt_cell.hyperlink = txt_file
    txt_cell.font = Font(color="0000FF", underline="single")
    
    folder_cell = ws.cell(row=i+1, column=3)
    folder_cell.value = batch_img_dir
    
    ws.cell(row=i+1, column=4, value=image_count)

ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 80
ws.column_dimensions['D'].width = 20

try:
    wb.save(output_excel)
    print("Excel updated!")
except PermissionError:
    print("PERMISSION ERROR: User has Excel open.")
