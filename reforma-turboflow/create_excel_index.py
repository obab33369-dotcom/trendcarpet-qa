import os
import shutil
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Paths
WORKSPACE_DIR = r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\reforma-turboflow"
ONEDRIVE_DIR = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures"
EXCEL_PATH = os.path.join(ONEDRIVE_DIR, "turboflow_batch_oversikt.xlsx")

def copy_missing_files():
    print("[INFO] Copying missing Batch 2 files to OneDrive...")
    # Copy CSV and JSON files for Batch 2 to OneDrive to keep it perfectly complete
    for part in (1, 2):
        for ext in ("csv", "json"):
            filename = f"turboflow_ready_batch2_del{part}.{ext}"
            src = os.path.join(WORKSPACE_DIR, filename)
            dest = os.path.join(ONEDRIVE_DIR, filename)
            if os.path.exists(src):
                shutil.copy2(src, dest)
                print(f"  -> Copied {filename} to OneDrive.")
            else:
                print(f"  -> Warning: {src} not found!")

def get_file_and_line_counts(batch_num, part_num):
    folder_name = f"turboflow_batch{batch_num}_produkter_aktiva_del{part_num}"
    folder_path = os.path.join(ONEDRIVE_DIR, folder_name)
    txt_name = f"turboflow_ready_batch{batch_num}_del{part_num}.txt"
    txt_path = os.path.join(ONEDRIVE_DIR, txt_name)
    
    file_count = 0
    if os.path.exists(folder_path):
        file_count = len([f for f in os.listdir(folder_path) if os.path.isfile(os.path.join(folder_path, f))])
        
    line_count = 0
    if os.path.exists(txt_path):
        with open(txt_path, "r", encoding="utf-8") as f:
            line_count = len([line for line in f if line.strip()])
            
    return file_count, line_count

def create_excel():
    print("[INFO] Creating premium Excel overview...")
    
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    # ----------------- Styles -----------------
    # Fonts
    font_title = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="2D3748")
    font_regular = Font(name="Segoe UI", size=10, color="2D3748")
    font_link = Font(name="Segoe UI", size=10, color="2B6CB0", underline="single")
    font_desc = Font(name="Segoe UI", size=9, italic=True, color="718096")
    
    # Fills
    fill_title = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Slate
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid") # Slate Blue
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Very light gray
    fill_accent = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Slate light
    
    # Alignments
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    # Borders
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    # Headers
    headers = ["Underbatch", "Unika Bildfiler i Mappen", "Antal Prompts", "Länk till Prompt-fil (.txt)", "Länk till Bildmapp (OneDrive)"]
    
    # ----------------- Sheet 1: Batch 1 -----------------
    ws1 = wb.create_sheet(title="Batch 1 (Möbler + Mattor)")
    ws1.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws1.merge_cells("A1:E1")
    ws1["A1"] = "REFORMA STOCKHOLM - BATCH 1 ÖVERSIKT"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_title
    ws1["A1"].alignment = align_center
    ws1.row_dimensions[1].height = 40
    
    # Description / Info block
    ws1["A3"] = "Info:"
    ws1["A3"].font = font_bold
    ws1["B3"] = "Batch 1 innehåller totalt 186 aktiva möbel-produkter samt samtliga 116 aktiva mattor från Reforma-katalogen."
    ws1["B3"].font = font_regular
    ws1["B4"] = "För optimal prestanda i Turboflow har prompts delats upp i 8 delar, var och en med tillhörande bildmapp."
    ws1["B4"].font = font_regular
    ws1["B5"] = "Skoningslös optimering: Varje mapp innehåller EXAKT de bildfiler (frön + mattor) som faktiskt refereras i just den delens prompts. Ingen onödig överföring!"
    ws1["B5"].font = font_desc
    
    ws1.row_dimensions[3].height = 18
    ws1.row_dimensions[4].height = 18
    ws1.row_dimensions[5].height = 18
    
    ws1.append([]) # row 6 blank
    
    row_idx = 7
    for col_idx, text in enumerate(headers, 1):
        cell = ws1.cell(row=row_idx, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all
    ws1.row_dimensions[row_idx].height = 28
    
    # Populate dynamic rows for Batch 1
    total_batch1_images = set()
    total_batch1_prompts = 0
    
    for p in range(1, 9):
        file_count, line_count = get_file_and_line_counts(1, p)
        total_batch1_prompts += line_count
        
        row_idx += 1
        ws1.row_dimensions[row_idx].height = 24
        
        # Sub-batch name
        c = ws1.cell(row=row_idx, column=1, value=f"Del {p}")
        c.font = font_bold
        c.alignment = align_center
        c.border = border_all
        
        # File Count
        c = ws1.cell(row=row_idx, column=2, value=file_count)
        c.font = font_regular
        c.alignment = align_center
        c.border = border_all
        
        # Prompts Count
        c = ws1.cell(row=row_idx, column=3, value=line_count)
        c.font = font_regular
        c.alignment = align_center
        c.border = border_all
        
        # Prompt File Hyperlink
        txt_file = f"turboflow_ready_batch1_del{p}.txt"
        abs_txt_path = os.path.abspath(os.path.join(ONEDRIVE_DIR, txt_file))
        formula_txt = f'=HYPERLINK("{abs_txt_path}", "{abs_txt_path}")'
        c = ws1.cell(row=row_idx, column=4, value=formula_txt)
        c.font = font_link
        c.alignment = align_left
        c.border = border_all
        
        # Image Folder Hyperlink
        img_folder = f"turboflow_batch1_produkter_aktiva_del{p}"
        abs_img_path = os.path.abspath(os.path.join(ONEDRIVE_DIR, img_folder))
        formula_img = f'=HYPERLINK("{abs_img_path}", "{abs_img_path}")'
        c = ws1.cell(row=row_idx, column=5, value=formula_img)
        c.font = font_link
        c.alignment = align_left
        c.border = border_all
        
        # Zebra shading
        if row_idx % 2 == 0:
            for col in range(1, 6):
                ws1.cell(row=row_idx, column=col).fill = fill_zebra
                
    # Add Total row for Batch 1
    row_idx += 1
    ws1.row_dimensions[row_idx].height = 24
    
    ws1.cell(row=row_idx, column=1, value="Totalt").font = font_bold
    ws1.cell(row=row_idx, column=1).alignment = align_center
    ws1.cell(row=row_idx, column=1).border = border_all
    ws1.cell(row=row_idx, column=1).fill = fill_accent
    
    ws1.cell(row=row_idx, column=2, value="302 unika prod").font = font_bold
    ws1.cell(row=row_idx, column=2).alignment = align_center
    ws1.cell(row=row_idx, column=2).border = border_all
    ws1.cell(row=row_idx, column=2).fill = fill_accent
    
    ws1.cell(row=row_idx, column=3, value=total_batch1_prompts).font = font_bold
    ws1.cell(row=row_idx, column=3).alignment = align_center
    ws1.cell(row=row_idx, column=3).border = border_all
    ws1.cell(row=row_idx, column=3).fill = fill_accent
    
    ws1.cell(row=row_idx, column=4, value="").border = border_all
    ws1.cell(row=row_idx, column=4).fill = fill_accent
    ws1.cell(row=row_idx, column=5, value="").border = border_all
    ws1.cell(row=row_idx, column=5).fill = fill_accent
    
    # ----------------- Sheet 2: Batch 2 -----------------
    ws2 = wb.create_sheet(title="Batch 2 (Möbler)")
    ws2.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws2.merge_cells("A1:E1")
    ws2["A1"] = "REFORMA STOCKHOLM - BATCH 2 ÖVERSIKT"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_title
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 40
    
    # Description / Info block
    ws2["A3"] = "Info:"
    ws2["A3"].font = font_bold
    ws2["B3"] = "Batch 2 innehåller totalt 173 aktiva möbel-produkter (del 2 av katalogen)."
    ws2["B3"].font = font_regular
    ws2["B4"] = "Batch 2 har delats upp i 2 delar med totalt 692 prompts."
    ws2["B4"].font = font_regular
    ws2["B5"] = "Även här har bildmapparna rensats extremt hårt för att bara innehålla de bilder som promptarna faktiskt refererar."
    ws2["B5"].font = font_desc
    
    ws2.row_dimensions[3].height = 18
    ws2.row_dimensions[4].height = 18
    ws2.row_dimensions[5].height = 18
    
    ws2.append([]) # row 6 blank
    
    row_idx = 7
    for col_idx, text in enumerate(headers, 1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all
    ws2.row_dimensions[row_idx].height = 28
    
    # Populate dynamic rows for Batch 2
    total_batch2_prompts = 0
    
    for p in range(1, 3):
        file_count, line_count = get_file_and_line_counts(2, p)
        total_batch2_prompts += line_count
        
        row_idx += 1
        ws2.row_dimensions[row_idx].height = 24
        
        # Sub-batch name
        c = ws2.cell(row=row_idx, column=1, value=f"Del {p}")
        c.font = font_bold
        c.alignment = align_center
        c.border = border_all
        
        # File Count
        c = ws2.cell(row=row_idx, column=2, value=file_count)
        c.font = font_regular
        c.alignment = align_center
        c.border = border_all
        
        # Prompts Count
        c = ws2.cell(row=row_idx, column=3, value=line_count)
        c.font = font_regular
        c.alignment = align_center
        c.border = border_all
        
        # Prompt File Hyperlink
        txt_file = f"turboflow_ready_batch2_del{p}.txt"
        abs_txt_path = os.path.abspath(os.path.join(ONEDRIVE_DIR, txt_file))
        formula_txt = f'=HYPERLINK("{abs_txt_path}", "{abs_txt_path}")'
        c = ws2.cell(row=row_idx, column=4, value=formula_txt)
        c.font = font_link
        c.alignment = align_left
        c.border = border_all
        
        # Image Folder Hyperlink
        img_folder = f"turboflow_batch2_produkter_aktiva_del{p}"
        abs_img_path = os.path.abspath(os.path.join(ONEDRIVE_DIR, img_folder))
        formula_img = f'=HYPERLINK("{abs_img_path}", "{abs_img_path}")'
        c = ws2.cell(row=row_idx, column=5, value=formula_img)
        c.font = font_link
        c.alignment = align_left
        c.border = border_all
        
        # Zebra shading
        if row_idx % 2 == 0:
            for col in range(1, 6):
                ws2.cell(row=row_idx, column=col).fill = fill_zebra
                
    # Add Total row for Batch 2
    row_idx += 1
    ws2.row_dimensions[row_idx].height = 24
    
    ws2.cell(row=row_idx, column=1, value="Totalt").font = font_bold
    ws2.cell(row=row_idx, column=1).alignment = align_center
    ws2.cell(row=row_idx, column=1).border = border_all
    ws2.cell(row=row_idx, column=1).fill = fill_accent
    
    ws2.cell(row=row_idx, column=2, value="173 unika prod").font = font_bold
    ws2.cell(row=row_idx, column=2).alignment = align_center
    ws2.cell(row=row_idx, column=2).border = border_all
    ws2.cell(row=row_idx, column=2).fill = fill_accent
    
    ws2.cell(row=row_idx, column=3, value=total_batch2_prompts).font = font_bold
    ws2.cell(row=row_idx, column=3).alignment = align_center
    ws2.cell(row=row_idx, column=3).border = border_all
    ws2.cell(row=row_idx, column=3).fill = fill_accent
    
    ws2.cell(row=row_idx, column=4, value="").border = border_all
    ws2.cell(row=row_idx, column=4).fill = fill_accent
    ws2.cell(row=row_idx, column=5, value="").border = border_all
    ws2.cell(row=row_idx, column=5).fill = fill_accent
    
    # ----------------- Auto-fit Column Widths on both sheets -----------------
    for ws in (ws1, ws2):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            
            for cell in col:
                val = str(cell.value or "")
                # Clean formulas for length check
                if val.startswith("="):
                    parts = val.split('"')
                    if len(parts) >= 4:
                        val = parts[3] # friendly name
                    elif len(parts) >= 2:
                        val = parts[1] # file path
                
                # Check line breaks
                for line in val.split("\n"):
                    if len(line) > max_len:
                        max_len = len(line)
            
            # Apply padding
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    # Save workbook
    try:
        wb.save(EXCEL_PATH)
        print(f"[SUCCESS] Dynamic Premium Excel created successfully at: {EXCEL_PATH}")
    except PermissionError:
        print(f"[ERROR] Permission denied when writing to: {EXCEL_PATH}")
        print("  -> DET VERKAR SOM ATT EXCEL-FILEN ÄR ÖPPEN PÅ DIN DATOR.")
        print("  -> Vänligen stäng 'turboflow_batch_oversikt.xlsx' i Excel och kör skriptet igen!")

if __name__ == "__main__":
    copy_missing_files()
    create_excel()
