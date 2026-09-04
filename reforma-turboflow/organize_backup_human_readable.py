import os
import sys
import json
import shutil
import re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Configure standard streams to UTF-8
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

BRAND_DICT_PATH = r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\reforma-turboflow\brand_sku_dict.json"
SOURCE_BACKUP_DIR = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\reforma_original_images_backup"
DEST_BY_PRODUCT_DIR = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\reforma_original_images_by_product"
EXCEL_OUTPUT_PATH = os.path.join(DEST_BY_PRODUCT_DIR, "Product_Catalog_Index.xlsx")

def clean_folder_name(name):
    # Remove characters that are invalid in Windows folder names
    invalid_chars = ['<', '>', ':', '"', '/', '\\', '|', '?', '*']
    for c in invalid_chars:
        name = name.replace(c, '')
    # Strip double spaces and surrounding quotes
    name = re.sub(r'\s+', ' ', name)
    return name.strip().strip("'").strip('"')

def main():
    print("==================================================")
    print("     ORGANIZING BACKUP INTO SUBFOLDER FTP MAPS    ")
    print("==================================================")
    
    if not os.path.exists(BRAND_DICT_PATH):
        print(f"❌ Error: Brand dictionary not found at {BRAND_DICT_PATH}!")
        return
    if not os.path.exists(SOURCE_BACKUP_DIR):
        print(f"❌ Error: Source backup directory not found at {SOURCE_BACKUP_DIR}!")
        return
        
    with open(BRAND_DICT_PATH, "r", encoding="utf-8") as f:
        brand_sku_dict = json.load(f)
        
    os.makedirs(DEST_BY_PRODUCT_DIR, exist_ok=True)
    
    # Establish source directories
    src_artiklar = os.path.join(SOURCE_BACKUP_DIR, "artiklar")
    src_liten = os.path.join(src_artiklar, "liten")
    src_zoom = os.path.join(src_artiklar, "zoom")
    
    print(f"Loaded {len(brand_sku_dict)} products from dictionary.")
    print("Copying and structuring files with individual FTP subdirectories...")
    
    excel_rows = []
    copied_count = 0
    
    for idx, (slug, info) in enumerate(brand_sku_dict.items()):
        sku = info["sku"]
        raw_name = info["name"]
        
        # Determine the file extension from the scraped image path
        rel_img_path = info["image_path"]
        _, ext = os.path.splitext(rel_img_path.split('?')[0])
        if not ext:
            ext = '.jpg'
            
        # Clean folder name: "Product Name (SKU)"
        clean_name = clean_folder_name(raw_name)
        folder_name = f"{clean_name} ({sku})"
        product_dest_dir = os.path.join(DEST_BY_PRODUCT_DIR, folder_name)
        
        # Create Askås-compliant subdirectories inside the product folder
        prod_artiklar_dir = os.path.join(product_dest_dir, "artiklar")
        prod_liten_dir = os.path.join(prod_artiklar_dir, "liten")
        prod_zoom_dir = os.path.join(prod_artiklar_dir, "zoom")
        
        os.makedirs(product_dest_dir, exist_ok=True)
        os.makedirs(prod_artiklar_dir, exist_ok=True)
        os.makedirs(prod_liten_dir, exist_ok=True)
        os.makedirs(prod_zoom_dir, exist_ok=True)
        
        # Search and copy files belonging to this product into subfolders
        copied_files = []
        
        # 1. Main Normal -> /artiklar/[SKU].jpg
        normal_src_file = os.path.join(src_artiklar, f"{sku}{ext}")
        if os.path.exists(normal_src_file):
            normal_dest_file = os.path.join(prod_artiklar_dir, f"{sku}{ext}")
            shutil.copy2(normal_src_file, normal_dest_file)
            copied_files.append(normal_dest_file)
            
        # 2. Thumbnail -> /artiklar/liten/[SKU]_S.jpg
        liten_src_file = os.path.join(src_liten, f"{sku}_S{ext}")
        if os.path.exists(liten_src_file):
            liten_dest_file = os.path.join(prod_liten_dir, f"{sku}_S{ext}")
            shutil.copy2(liten_src_file, liten_dest_file)
            copied_files.append(liten_dest_file)
            
        # 3. Zoom images (sequentially 1 to 10) -> /artiklar/zoom/[SKU]_[1-10].jpg
        for i in range(1, 11):
            zoom_src_file = os.path.join(src_zoom, f"{sku}_{i}{ext}")
            if os.path.exists(zoom_src_file):
                zoom_dest_file = os.path.join(prod_zoom_dir, f"{sku}_{i}{ext}")
                shutil.copy2(zoom_src_file, zoom_dest_file)
                copied_files.append(zoom_dest_file)
                
        copied_count += len(copied_files)
        
        # Record data for Excel (first image is in /artiklar/...)
        first_img_path = copied_files[0] if copied_files else ""
        excel_rows.append({
            "name": raw_name,
            "sku": sku,
            "url": info["url"],
            "first_image": first_img_path,
            "folder_path": product_dest_dir
        })
        
        if (idx + 1) % 100 == 0 or (idx + 1) == len(brand_sku_dict):
            print(f"  Progress: {idx + 1}/{len(brand_sku_dict)} products structured ({copied_count} files copied)")
            
    # Generate the highly styled Excel file
    print("\n📂 Generating master Excel index on OneDrive...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reforma Bildkatalog"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Styled headers
    headers = ["Artikelnamn", "SKU (Artikelnummer)", "Länk till Hemsida", "Lokal Förstabild", "Hyperlänk till Mapp"]
    ws.append(headers)
    
    # Styles config
    font_family = "Segoe UI"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark blue
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10)
    link_font = Font(name=font_family, size=10, color="0563C1", underline="single")
    
    zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Light gray zebra
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Format headers
    for col_num in range(1, 6):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    # Write data
    for idx, row in enumerate(sorted(excel_rows, key=lambda x: x["sku"])):
        row_num = idx + 2
        fill_color = zebra_fill if idx % 2 == 1 else white_fill
        
        # A: Product Name
        cell_a = ws.cell(row=row_num, column=1, value=row["name"])
        cell_a.font = data_font
        cell_a.fill = fill_color
        cell_a.border = thin_border
        cell_a.alignment = Alignment(vertical="center")
        
        # B: SKU
        cell_b = ws.cell(row=row_num, column=2, value=row["sku"])
        cell_b.font = data_font
        cell_b.fill = fill_color
        cell_b.border = thin_border
        cell_b.alignment = Alignment(horizontal="center", vertical="center")
        
        # C: Live URL Link
        cell_c = ws.cell(row=row_num, column=3)
        cell_c.value = f'=HYPERLINK("{row["url"]}", "Öppna Hemsida 🌐")'
        cell_c.font = link_font
        cell_c.fill = fill_color
        cell_c.border = thin_border
        cell_c.alignment = Alignment(horizontal="center", vertical="center")
        
        # D: First Image Link
        cell_d = ws.cell(row=row_num, column=4)
        if row["first_image"]:
            # Format as local file URI
            img_uri = "file:///" + row["first_image"].replace('\\', '/')
            cell_d.value = f'=HYPERLINK("{img_uri}", "Visa bild 🖼️")'
            cell_d.font = link_font
        else:
            cell_d.value = "Ingen bild"
            cell_d.font = data_font
        cell_d.fill = fill_color
        cell_d.border = thin_border
        cell_d.alignment = Alignment(horizontal="center", vertical="center")
        
        # E: Folder Link
        cell_e = ws.cell(row=row_num, column=5)
        folder_uri = "file:///" + row["folder_path"].replace('\\', '/')
        cell_e.value = f'=HYPERLINK("{folder_uri}", "Öppna Mapp 📂")'
        cell_e.font = link_font
        cell_e.fill = fill_color
        cell_e.border = thin_border
        cell_e.alignment = Alignment(horizontal="center", vertical="center")
        
    # Auto-adjust column widths
    ws.row_dimensions[1].height = 28
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if val.startswith('='):
                if "Hemsida" in val: max_len = max(max_len, 18)
                elif "Visa bild" in val: max_len = max(max_len, 15)
                elif "Mapp" in val: max_len = max(max_len, 15)
            else:
                max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Freeze the header row
    ws.freeze_panes = "A2"
    
    wb.save(EXCEL_OUTPUT_PATH)
    print(f"✓ Saved Beautiful Excel Catalog Index to: {EXCEL_OUTPUT_PATH}")
    
    print("\n==================================================")
    print("             ORGANIZATION COMPLETED!              ")
    print("==================================================")
    print(f"📂 Human-Readable Folder: {DEST_BY_PRODUCT_DIR}")
    print(f"📄 Excel Index:           {EXCEL_OUTPUT_PATH}")
    print("==================================================")

if __name__ == "__main__":
    main()
