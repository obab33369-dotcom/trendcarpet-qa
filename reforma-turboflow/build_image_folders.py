import os
import shutil
import csv
import glob
import openpyxl
from openpyxl.styles import Font
import json
import re

search_dirs = [
    r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\TEST TOPAZ",
    r"C:\Users\AndronikLindgren\reforma_automation\reforma_arkiv",
    r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\Projects\REFORMA\reforma-turboflow\test_furniture",
    r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\Projects\REFORMA\reforma-turboflow\new_rugs",
    r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\Projects\REFORMA\reforma-turboflow\test_rugs",
    r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\Projects\REFORMA\reforma-turboflow\downloaded_missing_images"
]

# Dynamically add all turboflow batch folders from OneDrive
onedrive_pictures_dir = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures"
if os.path.exists(onedrive_pictures_dir):
    for entry in os.listdir(onedrive_pictures_dir):
        if entry.lower().startswith("turboflow_batch"):
            full_path = os.path.join(onedrive_pictures_dir, entry)
            if os.path.isdir(full_path):
                search_dirs.append(full_path)

project_dir = r"C:\Users\AndronikLindgren\.gemini\antigravity\scratch\Projects\REFORMA\reforma-turboflow"

# Load sku_map.json to build reverse slug and sku lookups
sku_map_path = os.path.join(project_dir, "sku_map.json")
slug_to_filenames = {}
sku_to_filenames = {}
if os.path.exists(sku_map_path):
    try:
        with open(sku_map_path, 'r', encoding='utf-8') as f:
            sku_map = json.load(f)
        for filename, info in sku_map.items():
            slug = info.get('slug')
            sku = info.get('sku')
            if slug:
                slug_to_filenames.setdefault(slug.lower(), []).append(filename)
            if sku:
                sku_to_filenames.setdefault(sku.lower(), []).append(filename)
    except Exception as e:
        print(f"Warning: Failed to load sku_map.json: {e}")

print("Building image index...")
image_index = {}
for search_dir in search_dirs:
    if not os.path.exists(search_dir): continue
    for root, dirs, files in os.walk(search_dir):
        # Skip batch folders to prevent infinite loops or copying from ourselves
        if "batch" in root.lower() and "_images" in root.lower(): continue
        for file in files:
            if file.endswith(('.png', '.webp', '.jpg', '.jpeg')):
                if file not in image_index:
                    image_index[file] = os.path.join(root, file)

print(f"Found {len(image_index)} unique images.")
image_index_lower = {k.lower(): v for k, v in image_index.items()}

output_excel = os.path.join(project_dir, "turboflow_batches_index.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Batches"

headers = ["Batch Number", "Prompts TXT File (Click to Open)", "Folder Path (Copy this)", "Unique Images Count"]
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.value = header
    cell.font = Font(bold=True)

# Count how many CSV files were generated (only match numeric batches to get the full 185-image folders)
batch_files = glob.glob(os.path.join(project_dir, "turboflow_tracking_log_full_catalog_batch[0-9].csv")) + \
              glob.glob(os.path.join(project_dir, "turboflow_tracking_log_full_catalog_batch[0-9][0-9].csv"))

batch_names = []
for bf in batch_files:
    m = re.search(r"batch(\d+)\.csv$", bf)
    if m:
        batch_names.append(m.group(1))

def get_sort_key(s):
    match = re.match(r'(\d+)([a-z]?)', s)
    if match:
        return (int(match.group(1)), match.group(2))
    return (999, s)

batch_names = sorted(batch_names, key=get_sort_key)
print(f"Processing {len(batch_names)} batches...")

cleared_folders = set()

for idx_1, b_name in enumerate(batch_names, 1):
    txt_file = os.path.join(project_dir, f"turboflow_ready_full_catalog_batch{b_name}.txt")
    csv_file = os.path.join(project_dir, f"turboflow_tracking_log_full_catalog_batch{b_name}.csv")
    
    # Prepare target directory
    m_base = re.match(r"^(\d+)", b_name)
    base_num = m_base.group(1) if m_base else b_name
    batch_img_dir = os.path.join(project_dir, f"batch{base_num}_images")
    
    if batch_img_dir not in cleared_folders:
        if os.path.exists(batch_img_dir):
            shutil.rmtree(batch_img_dir)
        os.makedirs(batch_img_dir, exist_ok=True)
        cleared_folders.add(batch_img_dir)
        
    required_images = set()
    if os.path.exists(csv_file):
        with open(csv_file, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f, delimiter=';')
            for row in reader:
                refs = row.get('References Utilized', '')
                if not refs: continue
                for ref in refs.split(';'):
                    ref = ref.strip()
                    if ref: required_images.add(ref)
    
    copied_count = 0
    for img in required_images:
        match_src = image_index.get(img) or image_index_lower.get(img.lower())
        
        # 2. Try matching via slug / sku map from sku_map.json
        if not match_src:
            mapped_files = slug_to_filenames.get(img.lower()) or sku_to_filenames.get(img.lower())
            if mapped_files:
                for mapped_file in mapped_files:
                    if mapped_file in image_index:
                        match_src = image_index[mapped_file]
                        break
                    base_mapped = os.path.splitext(mapped_file)[0]
                    for ext in ['.png', '.webp', '.jpg', '.jpeg']:
                        alt_name = base_mapped + ext
                        if alt_name in image_index:
                            match_src = image_index[alt_name]
                            break
                    if match_src:
                        break
        
        # 3. Try substring/normalized fallback matching
        if not match_src:
            def normalize_local(s):
                s = s.lower()
                repl = {'ö': 'o', 'ä': 'a', 'å': 'a', 'é': 'e', 'Ǿ': 'o', 'ø': 'o', 'æ': 'a'}
                for c, r in repl.items(): s = s.replace(c, r)
                # also map common broken sequences from the CSV
                s = s.replace('fǾtlj', 'fatolj')
                s = s.replace('mssing', 'massing')
                s = s.replace('valnt', 'valnot')
                s = re.sub(r'^\d+[_-]', '', s)
                s = re.sub(r'\.(jpg|jpeg|png|webp)$', '', s)
                s = re.sub(r'[^a-z0-9]', '', s)
                s = re.sub(r'(26uwonder|wonder|26u)$', '', s)
                return s

            norm_img = normalize_local(img)
            
            # First try exact normalized match
            for k in image_index:
                if norm_img == normalize_local(k):
                    match_src = image_index[k]
                    break
            
            # If no exact match, try prefix match (e.g. required is 'stolpinnstol', image is 'stolpinnstol1')
            if not match_src:
                for k in image_index:
                    norm_k = normalize_local(k)
                    if norm_k.startswith(norm_img):
                        match_src = image_index[k]
                        break
        
        if match_src:
            # Copy under the physical name on disk so it keeps prefix and extension
            dst_name = os.path.basename(match_src)
            dst = os.path.join(batch_img_dir, dst_name)
            if not os.path.exists(dst):
                try:
                    shutil.copy2(match_src, dst)
                    copied_count += 1
                except:
                    pass
            else:
                copied_count += 1
            
    print(f"Batch {b_name}: copied {copied_count}/{len(required_images)} images.")
    
    ws.cell(row=idx_1+1, column=1, value=f"Batch {b_name}")
    
    txt_cell = ws.cell(row=idx_1+1, column=2)
    txt_cell.value = "Öppna TXT-fil"
    txt_cell.hyperlink = txt_file
    txt_cell.font = Font(color="0000FF", underline="single")
    
    folder_cell = ws.cell(row=idx_1+1, column=3)
    folder_cell.value = batch_img_dir
    
    ws.cell(row=idx_1+1, column=4, value=len(required_images))

ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 80
ws.column_dimensions['D'].width = 20

wb.save(output_excel)
print("Done!")
