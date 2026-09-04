import openpyxl
path = r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\26-06-05-in-out\Bade_IK100101936_Foto.xlsx"
wb = openpyxl.load_workbook(path)
print("Sheet names:", wb.sheetnames)
for name in wb.sheetnames:
    sheet = wb[name]
    print(f"Sheet {name}: {sheet.max_row} rows, {sheet.max_column} columns")
