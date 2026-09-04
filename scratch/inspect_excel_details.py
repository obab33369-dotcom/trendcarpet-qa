import openpyxl
import os

files = [
    r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\26-06-05-in-out\Bade_IK100101936_Foto.xlsx",
    r"C:\Users\AndronikLindgren\OneDrive - CaMa Gruppen AB\Pictures\26-06-12-Kids-rug\MAIN\Bade_IK100101936_Foto.xlsx"
]

for path in files:
    print(f"\n=========================================\nInspecting Excel file: {path}")
    if os.path.exists(path):
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            print("Sheets:", wb.sheetnames)
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                print(f"Sheet '{sheetname}': max_row={sheet.max_row}")
                
                # print first 5 rows
                rows = list(sheet.iter_rows(values_only=True))
                print("First 5 rows:")
                for r in rows[:5]:
                    print("  ", r)
        except Exception as e:
            print("Error reading excel:", e)
    else:
        print("File does not exist.")
